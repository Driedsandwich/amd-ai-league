"""手元（Mac・模型なし）の試験の部品。合成の束を tmp に作る。"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

APP = Path(__file__).resolve().parent.parent / "app"
sys.path.insert(0, str(APP))


def make_pdf(path: Path, lines):
    """依存なしで文字つきの 1 頁 PDF を手で書く。"""
    stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({l}) Tj T*" for l in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offs:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)


@pytest.fixture
def corpus(tmp_path):
    """読めない・暗号化・知らない型・空のフォルダ・空のファイルを含む合成の束。"""
    c = tmp_path / "corpus"
    for d in ("specs", "planning", "support", "logs", "engineering", "vendor", "archive"):
        (c / d).mkdir(parents=True)
    make_pdf(c / "specs" / "widget_datasheet_r2.pdf", ["Widget W9 maximum junction temperature is 94 C."])
    make_pdf(c / "specs" / "widget_datasheet_r1_WITHDRAWN.pdf",
             ["WITHDRAWN. Widget W9 maximum junction temperature is 105 C."])
    import docx
    d = docx.Document()
    d.add_paragraph("Roadmap FY27 overview paragraph.")
    t = d.add_table(rows=2, cols=3)
    for i, v in enumerate(["Product", "Milestone", "Quarter"]):
        t.rows[0].cells[i].text = v
    for i, v in enumerate(["W60", "customer sampling", "Q3 FY27"]):
        t.rows[1].cells[i].text = v
    d.save(str(c / "planning" / "roadmap_fy27.docx"))
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append(["part", "assembly", "number"])
    ws.append(["fan", "W9 field-replaceable fan", "ZZ-FAN-1-B"])
    wb.save(str(c / "support" / "rma_parts.xlsx"))
    (c / "support" / "bugs.csv").write_text(
        "ticket,summary,fixed_in\nZZ-100,Sensor high under load,4.3.2\nZZ-101,Hang,4.3.1\n", encoding="utf-8")
    (c / "logs" / "run.log").write_text(
        "2026-09-02 INFO ok\n2026-09-02 ERROR E7731: thermal throttle engaged\n2026-09-02 INFO ticket ZZ-100\n",
        encoding="utf-8")
    (c / "engineering" / "svc.py").write_text(
        '# old value 60\nDEFAULT_BATCH_TIMEOUT_S = 180\n', encoding="utf-8")
    (c / "vendor" / "capture.dat").write_bytes(bytes(range(256)) * 16)
    (c / "vendor" / "empty.txt").write_bytes(b"")
    (c / "vendor" / "audit.txt").write_text("secret audit text unreadable", encoding="utf-8")
    from pypdf import PdfWriter, PdfReader
    plain = tmp_path / "plain.pdf"
    make_pdf(plain, ["price is 6412"])
    w = PdfWriter(clone_from=str(plain))
    w.encrypt("pw", algorithm="RC4-128")
    with open(c / "vendor" / "supplier_ENCRYPTED.pdf", "wb") as f:
        w.write(f)
    (c / "support" / "label.jpg").write_bytes(b"\xff\xd8\xff\xe0" + b"\x00" * 64)
    return c


@pytest.fixture
def unreadable(corpus):
    """vendor/audit.txt を読めなくする（tmp の中だけ）。終わりに必ず戻す。"""
    p = corpus / "vendor" / "audit.txt"
    p.chmod(0o000)
    try:
        yield p
    finally:
        p.chmod(0o644)
