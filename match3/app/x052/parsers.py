"""型ごとの読み口。ファイルを読んで (場所, 文) の塊のリストを返す。

失敗は例外で知らせる。例外を捕まえて目録へ書くのは indexer の仕事。
返す一覧は Segments（list の子）で、上限・時間予算で打ち切ったときは truncated=True と理由を持つ。
deadline は time.monotonic() の絶対の時刻（None なら無期限）。
"""
from __future__ import annotations

import csv
import io
import os
import re
import time
from typing import List, Optional, Tuple

Segment = Tuple[str, str]  # (場所, 文)

CHUNK_SIZE = 800
CHUNK_OVERLAP = 150
ROW_MAX = 1500
TEXT_MAX_BYTES = 200 * 1024 * 1024   # 文・CSV の読む上限
XLSX_MAX_ROWS = 1_000_000            # 表 1 シートの読む行の上限

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
TEXT_EXTS = {
    ".txt", ".log", ".md", ".py", ".json", ".yaml", ".yml", ".ini", ".cfg", ".conf",
    ".toml", ".xml", ".html", ".htm", ".rst", ".tex", ".sh", ".js", ".ts", ".c", ".h",
    ".cpp", ".java", ".sql", ".jsonl", ".ndjson", ".tsv", ".env", ".properties",
}


class EncryptedError(Exception):
    pass


class UnknownTypeError(Exception):
    pass


class EmptyFileError(Exception):
    pass


class ImageFile(Exception):
    """画像。文字にするのは常駐の側。"""


class Segments(list):
    """塊の一覧。上限や時間で打ち切ったら truncated=True（理由は truncated_reason）。"""

    truncated = False
    truncated_reason = ""

    def mark(self, reason: str) -> None:
        if not self.truncated:
            self.truncated, self.truncated_reason = True, reason


def _expired(deadline: Optional[float]) -> bool:
    return deadline is not None and time.monotonic() > deadline


def chunk_plain(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]
    out, start, n = [], 0, len(text)
    while start < n:
        end = min(n, start + size)
        if end < n:
            cut = max(text.rfind("\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if cut > start:
                end = cut + 1
        piece = text[start:end].strip()
        if piece:
            out.append(piece)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return out


def chunk_lines(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP,
                deadline: Optional[float] = None) -> Segments:
    """行の境目を尊重して詰める。場所は 'lines a-b'。"""
    lines = text.splitlines()
    segs = Segments()
    buf: List[Tuple[int, str]] = []
    cur = 0

    def flush():
        if not buf:
            return
        body = "\n".join(t for _, t in buf).strip()
        if body:
            segs.append((f"lines {buf[0][0]}-{buf[-1][0]}", body))

    for no, line in enumerate(lines, 1):
        if no % 20000 == 0 and _expired(deadline):
            segs.mark("time")
            break
        if len(line) > size:  # 長すぎる 1 行は文字で割る
            flush()
            buf, cur = [], 0
            for i, piece in enumerate(chunk_plain(line, size, overlap)):
                segs.append((f"line {no} part {i + 1}", piece))
            continue
        if cur + len(line) + 1 > size and buf:
            flush()
            keep, tot = [], 0
            for item in reversed(buf):  # 重なり: 末尾の数行
                if tot + len(item[1]) > overlap:
                    break
                keep.insert(0, item)
                tot += len(item[1]) + 1
            buf, cur = keep, tot
        buf.append((no, line))
        cur += len(line) + 1
    flush()
    return segs


def _row_segments(loc: str, row_text: str) -> List[Segment]:
    if len(row_text) <= ROW_MAX:
        return [(loc, row_text)]
    return [(f"{loc} part {i + 1}", p) for i, p in enumerate(chunk_plain(row_text, CHUNK_SIZE, CHUNK_OVERLAP))]


def _read_capped(path: str) -> Tuple[bytes, bool]:
    """先頭から TEXT_MAX_BYTES まで読む。(中身, 上限で切ったか)"""
    with open(path, "rb") as f:
        data = f.read(TEXT_MAX_BYTES + 1)
    if len(data) > TEXT_MAX_BYTES:
        return data[:TEXT_MAX_BYTES], True
    return data, False


def decode_text(data: bytes) -> str:
    if data[:3] == b"\xef\xbb\xbf":
        data = data[3:]
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    return data.decode("utf-8", errors="replace")


def printable_ratio(data: bytes) -> float:
    if not data:
        return 0.0
    s = data.decode("utf-8", errors="replace")
    ok = sum(1 for ch in s if (ch.isprintable() and ch != "�") or ch in "\t\n\r")
    return ok / len(s)


def parse_pdf(path: str, deadline: Optional[float] = None) -> Segments:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
    except Exception:
        # 暗号化の PDF は開く時に例外になることがある（AES は cryptography が要る）。開かない。
        if b"/Encrypt" in _read_capped(path)[0]:
            raise EncryptedError("encrypted PDF")
        raise
    if reader.is_encrypted:
        raise EncryptedError("encrypted PDF")
    segs = Segments()
    for pno, page in enumerate(reader.pages, 1):
        if _expired(deadline):
            segs.mark("time")
            break
        try:
            text = page.extract_text() or ""
        except Exception:
            continue  # 頁 1 枚の失敗でファイルを捨てない
        for i, piece in enumerate(chunk_plain(text)):
            segs.append((f"p{pno}" if i == 0 else f"p{pno} part {i + 1}", piece))
    if not segs and not segs.truncated:
        raise EmptyFileError("PDF has no extractable text")
    return segs


def parse_docx(path: str, deadline: Optional[float] = None) -> Segments:
    import docx

    d = docx.Document(path)
    segs = Segments()
    paras = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    for i, piece in enumerate(chunk_plain("\n".join(paras))):
        segs.append((f"paragraphs part {i + 1}", piece))
    for tno, table in enumerate(d.tables, 1):
        rows = []
        for r in table.rows:
            if _expired(deadline):
                segs.mark("time")
                break
            rows.append([(c.text or "").strip() for c in r.cells])
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        header = rows[0]
        if len(rows) == 1:
            segs.extend(_row_segments(f"table {tno} row 1", " | ".join(x for x in rows[0] if x)))
            continue
        for rno, r in enumerate(rows[1:], 2):
            cells = [f"{h}={v}" if h else v for h, v in zip(header, r) if v]
            segs.extend(_row_segments(f"table {tno} row {rno}", " | ".join(cells)))
    if not segs and not segs.truncated:
        raise EmptyFileError("docx has no text")
    return segs


def parse_xlsx(path: str, deadline: Optional[float] = None) -> Segments:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    segs = Segments()
    try:
        for ws in wb.worksheets:
            header = None
            for rno, row in enumerate(ws.iter_rows(values_only=True), 1):
                if rno > XLSX_MAX_ROWS:
                    segs.mark(f"rows>{XLSX_MAX_ROWS}")
                    break
                if rno % 5000 == 0 and _expired(deadline):
                    segs.mark("time")
                    break
                vals = ["" if v is None else str(v).strip() for v in row]
                if not any(vals):
                    continue
                if header is None:
                    header = vals
                    segs.append((f"{ws.title}!row {rno}", f"{ws.title} | " + " | ".join(v for v in vals if v)))
                    continue
                cells = [f"{h}={v}" if h else v for h, v in zip(header, vals) if v]
                segs.extend(_row_segments(f"{ws.title}!row {rno}", f"{ws.title} | " + " | ".join(cells)))
            if segs.truncated_reason == "time":
                break
    finally:
        wb.close()
    if not segs and not segs.truncated:
        raise EmptyFileError("xlsx has no cells")
    return segs


def parse_csv(path: str, deadline: Optional[float] = None) -> Segments:
    data, cut = _read_capped(path)
    text = decode_text(data)
    if not text.strip():
        raise EmptyFileError("empty csv")
    csv.field_size_limit(10 * 1024 * 1024)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    segs = Segments()
    if cut:
        segs.mark(f"bytes>{TEXT_MAX_BYTES}")
    header = None
    for rno, r in enumerate(csv.reader(io.StringIO(text), dialect), 1):
        if rno % 20000 == 0 and _expired(deadline):
            segs.mark("time")
            break
        if not any(c.strip() for c in r):
            continue
        if header is None:
            header = [c.strip() for c in r]
            segs.append(("header", " | ".join(h for h in header if h)))
            continue
        cells = [f"{h}={v.strip()}" if h else v.strip() for h, v in zip(header, r) if v.strip()]
        extra = [v.strip() for v in r[len(header):] if v.strip()]
        segs.extend(_row_segments(f"row {rno}", " | ".join(cells + extra)))
    if not segs:
        raise EmptyFileError("empty csv")
    return segs


def parse_text(path: str, deadline: Optional[float] = None) -> Segments:
    data, cut = _read_capped(path)
    text = decode_text(data)
    if not text.strip():
        raise EmptyFileError("empty text")
    segs = chunk_lines(text, deadline=deadline)
    if cut:
        segs.mark(f"bytes>{TEXT_MAX_BYTES}")
    return segs


def kind_of(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext == ".csv":
        return "csv"
    if ext in IMAGE_EXTS:
        return "image"
    if ext in TEXT_EXTS:
        return "text"
    return "other"


def parse_file(path: str, deadline: Optional[float] = None) -> Segments:
    """読めた塊を返す。画像は ImageFile、読めない・暗号化・空・知らない型は例外。"""
    if os.path.getsize(path) == 0:
        raise EmptyFileError("empty file")
    kind = kind_of(path)
    if kind == "pdf":
        return parse_pdf(path, deadline)
    if kind == "docx":
        return parse_docx(path, deadline)
    if kind == "xlsx":
        return parse_xlsx(path, deadline)
    if kind == "csv":
        return parse_csv(path, deadline)
    if kind == "text":
        return parse_text(path, deadline)
    if kind == "image":
        with open(path, "rb") as f:  # 読めるかだけ確かめる（PermissionError をここで出す）
            f.read(16)
        raise ImageFile(os.path.splitext(path)[1].lower())
    # それ以外: 先頭 4KB の印字率で文か知らない型かを決める
    with open(path, "rb") as f:
        head = f.read(4096)
    if printable_ratio(head) >= 0.9:
        return parse_text(path, deadline)
    raise UnknownTypeError("unknown binary type")
